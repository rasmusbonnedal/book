import argparse
import datetime
import json
import os
from pathlib import Path

from borb.pdf import (
    Document,
    FixedColumnWidthTable,
    FlexibleColumnWidthTable,
    Heading,
    Image,
    LayoutElement,
    Page,
    Paragraph,
    PDF,
    SingleColumnLayout,
)
import borb.pdf
import jinja2
from openpyxl import load_workbook

company = {
    "name": "AB Schmaktiebolaget",
    "org_nr": "165560269986",
    "contact_name": "Alf Orm",
    "contact_tel": "0733123456",
    "contact_email": "alf.orm@schmaktie.se",
    "address": "Bofinkstigen 33",
    "ort": "Gnesta",
}

employee = {"name": "Ben Burt", "id": "198612312382", "account": "5169-0012345"}


def read_file(input_file):
    try:
        workbook = load_workbook(input_file, data_only=True, read_only=True)
    except PermissionError:
        raise SystemExit(
            f"Error: Cannot open '{input_file}'. Make sure it's not open in Excel."
        )
    sheet = workbook["MyrPay"]
    months = list(sheet.iter_rows(min_row=7, max_row=18, max_col=10, values_only=True))
    vab = sheet.iter_rows(min_row=21, max_col=2, values_only=True)
    vab = [x for x in vab if x[0] is not None]
    return (months, vab)


def skv_avrunda_nedat(belopp):
    return int(belopp)


def is_last(iter):
    l = len(iter)
    for i in range(l):
        yield iter[i], i == (l - 1)


def format_kr(belopp):
    s = f"{belopp:.2f} kr"
    x = 9
    while True:
        a = len(s) - x
        if a <= 0:
            break
        s = s[:a] + " " + s[a:]
        x += 4
    return s


def manads_namn(month):
    return [
        "januari",
        "februari",
        "mars",
        "april",
        "maj",
        "juni",
        "juli",
        "augusti",
        "september",
        "oktober",
        "november",
        "december",
    ][month - 1]


# month to generate report for [1, 12]
def report(months, vab, year, month, company_data, employee_data):
    if month < 1 or month > 12:
        raise Exception("Invalid month")
    m = months[month - 1]
    vab_month = month - 1 if month > 1 else 12
    current_vab = [x for x in vab if x[0].month == vab_month]

    month_name = m[0]
    brutto_lon = m[1]
    avdrag_vab = m[2]
    avdrag_skatt = skv_avrunda_nedat(m[3])
    bilforman = m[4] + m[5]
    nettolon_avdrag = m[5]
    traktamente = m[6] != None
    milersattning = m[7] != None
    nettolon_avdrag_ovr_60 = m[8] or 0.0
    nettolon_avdrag_ovr_100 = m[9] or 0.0
    brutto_kontant = brutto_lon + avdrag_vab
    print(brutto_lon)
    print(avdrag_vab)
    print(avdrag_skatt)
    print(nettolon_avdrag)
    print(nettolon_avdrag_ovr_100)
    print(nettolon_avdrag_ovr_60)
    utbetala = (
        brutto_lon
        + avdrag_vab
        + avdrag_skatt
        + nettolon_avdrag
        + nettolon_avdrag_ovr_60
        + nettolon_avdrag_ovr_100
    )

    hours = sum([y for _, y in current_vab])
    dates = ", ".join([f"{x.day}/{x.month}" for x, y in current_vab])
    if (hours == 0) ^ (avdrag_vab == 0):
        raise Exception(f"Error: vab-timmar: {hours}, vab-avdrag: {avdrag_vab}")
    if hours == 0:
        vab_spec = ""
    else:
        vab_spec = f" {hours} h ({dates})"

    borb.pdf.license.usage_statistics.UsageStatistics.opt_out()  # type: ignore
    doc = Document()
    page = Page()
    doc.append_page(page)
    layout = SingleColumnLayout(page)
    layout.append_layout_element(
        FixedColumnWidthTable(1, 2, column_widths=[4, 1])
        .append_layout_element(
            Heading(
                f"Lönespecifikation {company_data['name']}",
                font="Helvetica-bold",  # type: ignore
                font_size=16,
                vertical_alignment=LayoutElement.VerticalAlignment.TOP,
            )
        )
        .append_layout_element(
            Image(
                Path(company_data["logo"]),
                size=(64, 64),
                horizontal_alignment=LayoutElement.HorizontalAlignment.RIGHT,
            )
        )
        .no_borders()
    )

    table = FlexibleColumnWidthTable(3, 2)
    table = (
        table.append_layout_element(
            Paragraph("Namn", padding_right=20, font="Helvetica-bold")
        )
        .append_layout_element(Paragraph(employee_data["name"]))
        .append_layout_element(
            Paragraph("Bankkonto", padding_right=20, font="Helvetica-bold")
        )
        .append_layout_element(Paragraph(employee_data["account"]))
        .append_layout_element(
            Paragraph("Period", padding_right=20, font="Helvetica-bold")
        )
        .append_layout_element(Paragraph(f"{year:04d}-{month:02d}"))
        .no_borders()
    )
    layout.append_layout_element(table)

    layout.append_layout_element(Paragraph("", padding_top=30))

    table_data = []
    table_data.append((f"Månadslön {manads_namn(month)}", format_kr(brutto_lon)))
    if avdrag_vab != 0:
        table_data.append((f"Avdrag VAB{vab_spec}", format_kr(avdrag_vab)))
        table_data.append((f"Summa bruttolön", format_kr(brutto_kontant)))
    table_data.append((f"Avdragen skatt", format_kr(avdrag_skatt)))
    if nettolon_avdrag != 0:
        table_data.append((f"Nettolöneavdrag bilförmån", format_kr(nettolon_avdrag)))
    if nettolon_avdrag_ovr_60 != 0:
        table_data.append(
            (
                f"Nettolöneavdrag sjukvårdsförsäkring 60%",
                format_kr(nettolon_avdrag_ovr_60),
            )
        )
    if nettolon_avdrag_ovr_100 != 0:
        table_data.append(
            (
                f"Nettolöneavdrag sjukvårdsförsäkring närstående 100%",
                format_kr(nettolon_avdrag_ovr_100),
            )
        )
    if traktamente or milersattning:
        if traktamente and milersattning:
            spec = "Traktamente och milersättning"
        elif traktamente:
            spec = "Traktamente"
        else:
            spec = "Milersättning"
        table_data.append((f"{spec} har betalats ut under månaden", ""))
    table_data.append((f"Att utbetala", format_kr(utbetala)))

    table = FixedColumnWidthTable(len(table_data), 2, column_widths=[4, 1]).striped()
    first = True
    for d, isl in is_last(table_data):
        padding_top = 80 if isl else 5 if first else 0
        padding_bottom = 5 if isl else 0
        padding_side = 10
        first = False
        table.append_layout_element(
            Paragraph(
                d[0],
                padding_left=padding_side,
                padding_top=padding_top,
                padding_bottom=padding_bottom,
            )
        )
        table.append_layout_element(
            Paragraph(
                d[1],
                horizontal_alignment=LayoutElement.HorizontalAlignment.RIGHT,
                padding_right=padding_side,
                padding_top=padding_top,
                padding_bottom=padding_bottom,
            )
        )
    layout.append_layout_element(table.no_internal_borders())

    pdf_path = rf"C:\Users\rasmu\OneDrive\Documents\Myrkotten AB\Verifikat {year}\{datetime.date.today().strftime('%Y-%m-%d')} Lön.pdf"
    xml_path = rf"C:\Users\rasmu\OneDrive\Documents\Myrkotten AB\Bokslut {year}\agd\arbetsgivardeklaration-{manads_namn(month)}-{year}.xml"
    if os.path.isfile(pdf_path) or os.path.isfile(xml_path):
        existing = pdf_path if os.path.isfile(pdf_path) else xml_path
        raise SystemExit(
            f"Error: Output file '{existing}' already exists. Will not overwrite."
        )
    PDF.write(what=doc, where_to=pdf_path)

    if avdrag_skatt > 0:
        raise Exception("Avdragen skatt ska vara negativ")

    franvaro_uppgifter = []
    for vab in current_vab:
        franvaro_uppgifter.append(
            {
                "BetalningsmottagarId": employee_data["id"],
                "FranvaroDatum": vab[0].strftime("%Y-%m-%d"),
                "FranvaroProcentTFP": int(vab[1] * 100 / 8),
                "FranvaroTimmarTFP": vab[1],
            }
        )

    skv_brutto_kontant = skv_avrunda_nedat(brutto_kontant)
    skv_avdrprelskatt = skv_avrunda_nedat(-avdrag_skatt)
    iu = {
        "BetalningsmottagarId": employee_data["id"],
        "AvdrPrelSkatt": skv_avdrprelskatt,
        "KontantErsattningUlagAG": skv_brutto_kontant,
        "SkatteplOvrigaFormanerUlagAG": 0,
        "SkatteplBilformanUlagAG": skv_avrunda_nedat(bilforman),
        "DrivmVidBilformanUlagAG": 0,
        "AndraKostnadsers": 0,
        "BilErsattning": milersattning,
        "Traktamente": traktamente,
        "ArbetsplatsensGatuadress": company_data["address"],
        "ArbetsplatsensOrt": company_data["ort"],
        "franvaro_uppgifter": franvaro_uppgifter,
    }

    data = {
        "org_nummer": company_data["org_nr"],
        "kontakt_namn": company_data["contact_name"],
        "kontakt_tel_nummer": company_data["contact_tel"],
        "kontakt_email": company_data["contact_email"],
        "skapad_tid": datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        "period": f"{year:04d}{month:02d}",
        "total_arbavg": skv_avrunda_nedat(skv_brutto_kontant * 0.3142),
        "total_skatteavdr": skv_avdrprelskatt,
        "ius": [iu],
    }

    env = jinja2.Environment(
        loader=jinja2.FileSystemLoader("."), undefined=jinja2.StrictUndefined
    )
    template = env.get_template("agd.jinja", globals=data)
    with open(xml_path, "w", encoding="utf-8") as f:
        print(template.render(), file=f)


input_file = r"C:\Users\rasmu\OneDrive\Documents\Myrkotten AB\Lön och Utlägg 2026.xlsx"
months, vab = read_file(input_file)

if os.path.isfile("data.json"):
    with open("data.json") as f:
        data = json.load(f)
else:
    with open("demo.json") as f:
        data = json.load(f)

parser = argparse.ArgumentParser()
parser.add_argument("year", type=int, help="Year")
parser.add_argument("month", type=int, help="Month (1-12)")
args = parser.parse_args()

report(months, vab, args.year, args.month, data["company"], data["employee"])
